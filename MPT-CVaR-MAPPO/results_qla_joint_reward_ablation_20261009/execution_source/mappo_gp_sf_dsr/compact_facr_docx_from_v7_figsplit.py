"""Create a compact manuscript version from paper_draft_facr_cn_v7_fig1split.docx.

This script keeps the user's v7 DOCX as the source of truth, preserves existing
equations/images where possible, and removes research-log style material from
the main narrative.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "paper_draft_facr_cn_v7_fig1split.docx"
OUT = ROOT / "paper_draft_facr_cn_v8_compact.docx"


def set_para_text(para, text: str, style: str | None = None) -> None:
    para.clear()
    run = para.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "SimSun")
    if style:
        para.style = style


def remove_element(el) -> None:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def move_after(anchor_el, elements) -> None:
    """Move elements after anchor, preserving their order."""
    current = anchor_el
    for el in elements:
        current.addnext(el)
        current = el


def renumber_figures(doc: Document) -> None:
    # First collect surviving caption labels in document order.
    captions = []
    cap_re = re.compile(r"^Fig\.\s*(\d+[a-z]?)\.")
    for para in doc.paragraphs:
        text = para.text.strip()
        m = cap_re.match(text)
        if m:
            captions.append((para, m.group(1)))

    mapping: dict[str, str] = {}
    next_id = 1
    for _, old in captions:
        if old not in mapping:
            mapping[old] = str(next_id)
            next_id += 1

    # Replace longer keys first so "16a" is handled before "16".
    keys = sorted(mapping.keys(), key=len, reverse=True)
    ref_re = re.compile(r"\bFig\.\s*(" + "|".join(re.escape(k) for k in keys) + r")\b") if keys else None

    if ref_re:
        for para in doc.paragraphs:
            for run in para.runs:
                if "Fig." in run.text:
                    run.text = ref_re.sub(lambda m: f"Fig. {mapping[m.group(1)]}", run.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if "Fig." in run.text:
                                run.text = ref_re.sub(lambda m: f"Fig. {mapping[m.group(1)]}", run.text)


def main() -> None:
    doc = Document(str(SRC))
    paras = list(doc.paragraphs)
    tables = list(doc.tables)

    # Related work: compress seven subsections into three.
    set_para_text(paras[20], "2.1 MEC 任务卸载与学习型决策", "Heading 3")
    set_para_text(
        paras[38],
        "近期 MEC 研究还提出了 QoE-oriented distributed DRL、federated DRL、auction/MAB 和 incomplete-information offloading 等学习型方法 [21]–[23]。这些方向分别覆盖 QoE 驱动卸载、队列约束学习和不完整信息下的探索利用，但其原始动作空间、信息边界和系统接口通常与本文的 stale-summary portfolio offloading 不完全一致。因此，本文在实验中将其作为 adapted baselines，在统一动作接口、反馈信息、训练预算和评估环境下比较其核心决策思想。",
        "Body Text",
    )
    # Move the concise learned-baseline paragraph under 2.1.
    paras[21]._p.addnext(paras[38]._p)
    for idx in [37, 39, 40]:
        remove_element(paras[idx]._p)

    set_para_text(paras[22], "2.2 陈旧负载信息与反馈新鲜度", "Heading 3")
    set_para_text(
        paras[34],
        "AoI/AoT 相关研究说明，MEC 决策不仅取决于信息内容，也取决于信息的新鲜度 [14]–[17], [29]。与这些工作主要关注任务数据、监测结果或计算结果的新鲜度不同，本文关注用户做卸载决策时使用的服务器负载摘要是否新鲜。二者共同说明信息年龄具有决策意义，但本文进一步研究服务器摘要年龄如何通过多用户同步 portfolio response 诱发系统性拥挤。",
        "Body Text",
    )
    paras[24]._p.addnext(paras[34]._p)
    for idx in [33, 35, 36]:
        remove_element(paras[idx]._p)

    set_para_text(paras[25], "2.3 尾部风险与队列感知卸载", "Heading 3")
    for idx in [29]:
        remove_element(paras[idx]._p)
    set_para_text(
        paras[31],
        "这些 queue-aware 方法为 FACR 的 stale Lyapunov prior 提供了理论基础：卸载决策应响应队列压力，而不能只追求瞬时传输或服务代价。本文的关键区别是，在线 Actor 不访问 fresh queue，只能使用 B(t−δ) 中的滞后摘要构造低置信度结构参考。因此，正式 baseline 不使用 fresh queue oracle；消融中的 Prior-only 则用于判断同一 stale summary 和同一 portfolio fusion 约束下，纯 prior 是否已经足够。",
        "Body Text",
    )
    remove_element(paras[32]._p)
    remove_element(paras[41]._p)
    remove_element(tables[0]._element)
    remove_element(paras[42]._p)

    # System model: remove standalone symbol-section style, fold mechanism insight into Chapter 3.
    set_para_text(paras[63], "Table 1. List of Commonly Used Symbols.", "Body Text")
    paras[63].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_text(paras[64], "3.4 假设、动作映射与技术挑战", "Heading 3")
    set_para_text(paras[72], "3.5 局部机理分析", "Heading 3")
    remove_element(paras[73]._p)

    # Move the uniform load-regime scan into the experiments chapter and remove the old chapter shell.
    moved_load_blocks = [paras[78]._p, paras[79]._p, paras[80]._p, paras[83]._p]
    remove_element(paras[77]._p)
    remove_element(paras[81]._p)
    remove_element(paras[82]._p)
    remove_element(paras[84]._p)

    # Method chapter renumbering and pruning of implementation-log prose.
    set_para_text(paras[85], "4 FACR-MAPPO 方法", "Heading 2")
    set_para_text(paras[86], "4.1 总体架构与在线 portfolio 构造", "Heading 3")
    set_para_text(paras[90], "Feedback-age-aware Actor proposal.", "Body Text")
    set_para_text(paras[106], "Structured stale Lyapunov prior.", "Body Text")
    set_para_text(paras[115], "Prior-guided stabilized portfolio fusion.", "Body Text")
    set_para_text(paras[124], "4.2 CTDE 与风险感知 PPO 目标", "Heading 3")
    set_para_text(paras[138], "4.3 训练算法、在线执行与复杂度", "Heading 3")
    set_para_text(paras[141], "复杂度与实现说明。", "Body Text")
    remove_element(paras[145]._p)

    # Experiments chapter: compact hypotheses and parameter-screening material.
    set_para_text(paras[146], "5 实验评估", "Heading 2")
    set_para_text(paras[147], "5.1 实验设置、baseline 与指标", "Heading 3")
    set_para_text(
        paras[149],
        "实验围绕四个假设展开：H1 验证 stale feedback 是否在临界高利用率区显著恶化 tail latency 和 queue oscillation；H2 验证 FACR 是否优于风险感知、队列控制和不完整信息 baselines；H3 验证 Lyapunov prior、AgeGate 与学习型 Actor 的职责边界；H4 验证固定 checkpoint 在跨轨迹、请求分布和用户规模变化下的泛化能力，以及在线推理开销是否低于 slot 时长。",
        "Body Text",
    )
    for idx in [150, 151, 152, 153]:
        remove_element(paras[idx]._p)
    set_para_text(
        paras[154],
        "训练预算采用预定义的 10k steps，并使用相同 seed、feedback-age 网格和测试长度比较各方法。50k 延长训练仅作为 checkpoint 选择和训练稳定性诊断，不作为主要有效性结果；跨环境鲁棒性实验采用“不重训”的固定 checkpoint 协议。",
        "Body Text",
    )
    set_para_text(
        paras[155],
        "参数配置原则。FACR-MAPPO 的候选范围由参数语义和相关文献确定；分式析因用于筛选敏感参数，局部扫描用于选择最终折中。正文仅报告最终网络与关键超参数，完整筛选曲线、主效应热图和候选组合结果放入补充材料。",
        "Body Text",
    )
    # Remove local validation table/figures and detailed screening figures.
    remove_element(tables[5]._element)
    for idx in [158, 159, 160, 161, 162, 163, 164, 165, 166, 167, 168, 169]:
        remove_element(paras[idx]._p)

    # Training dynamics moved out of main text.
    remove_element(paras[171]._p)
    remove_element(paras[172]._p)
    remove_element(tables[7]._element)
    remove_element(paras[173]._p)
    remove_element(paras[174]._p)

    # Problem existence / load regimes.
    set_para_text(paras[175], "5.2 问题存在性与负载区间", "Heading 3")
    set_para_text(
        paras[176],
        "本节只回答一个问题：stale feedback 在什么负载区域产生系统性拥挤。机制参照包括 score-softmax、delayed-JSQ 和 FACR-MAPPO。score-softmax 根据滞后评分选择节点，delayed-JSQ 追逐滞后最短队列；二者都容易在状态更新后产生同步响应，而 FACR 通过 AgeGate、stale prior 和保守 portfolio fusion 降低对过期摘要的过度响应。",
        "Body Text",
    )
    move_after(paras[176]._p, moved_load_blocks)
    remove_element(paras[177]._p)
    remove_element(paras[178]._p)
    remove_element(paras[179]._p)
    remove_element(paras[180]._p)
    remove_element(paras[181]._p)

    # Main baseline comparison.
    set_para_text(paras[182], "5.3 主 baseline 结果", "Heading 3")
    set_para_text(
        paras[183],
        "Baseline 命名中，A 表示 adapted，即在保留核心决策机制的前提下适配到本文统一的 MEC portfolio/stale-summary 环境；LC 表示 local-candidate Lyapunov-control adaptation。Adapted baseline 不等于原论文完整系统逐项复现，而是在统一动作语义、反馈信息、训练预算和评价环境后的可比实现。",
        "Body Text",
    )
    remove_element(paras[185]._p)
    remove_element(paras[186]._p)

    # Core ablation: keep focused variants and AgeGate sensitivity, remove detailed diagnostics from main text.
    set_para_text(paras[189], "5.4 核心消融与结构诊断", "Heading 3")
    set_para_text(
        paras[191],
        "核心消融汇总 KAIST、N=30、poisson-bursty 请求、δ∈{0,4,8} 和 3 个训练 seed 下的结果。这里的 migration 表示相邻时隙 portfolio action 的平均变化幅度，即用户平均的 ||w_i(t)−w_i(t−1)||_1；它衡量卸载权重是否频繁迁移，而不是用户物理移动。",
        "Body Text",
    )
    remove_element(paras[194]._p)
    remove_element(paras[195]._p)
    set_para_text(
        paras[196],
        "消融汇总表明，去掉 Lyapunov prior 后，CVaR99 从 3.539 s 上升到 4.652 s，恶化约 31.5%；平均时延和队列振荡也分别从 0.926 s、0.451 上升到 0.979 s、0.502。这说明 stale Lyapunov prior 是 FACR-MAPPO 的主要结构化负载控制来源。Actor-only MAPPO 的 CVaR99 为 4.454 s，也明显弱于完整 FACR，说明普通共享 MAPPO Actor 在 stale server summary 场景中难以单独稳定控制尾部风险。",
        "Body Text",
    )
    set_para_text(
        paras[197],
        "AgeGate 的作用主要体现在反馈年龄敏感性，而不是总体均值指标。w/o AgeGate 的平均时延为 0.924 s，与完整 FACR 的 0.926 s 基本相同；CVaR99 也仅从 3.539 s 变为 3.552 s。因此，本文不把 AgeGate 表述为主要平均性能来源，而将其定位为 feedback-age robustness module。定义",
        "Body Text",
    )
    set_para_text(
        paras[203],
        "Prior-only 的整体 CVaR99 为 3.546 s，几乎等于完整 FACR 的 3.539 s，说明 stale Lyapunov prior 本身已经是强控制器。但 Prior-only 的 migration 为 0.0204，约为完整 FACR 的 3.45 倍；同时其 CVaR99 随 feedback delay 明显退化，δ=0、4、8 下分别为 3.025 s、3.611 s 和 3.979 s。由此可见，prior 在新鲜或轻度滞后状态下非常有效，但直接追随 stale queue summary 会造成更大的 portfolio 迁移和更强的反馈年龄退化。",
        "Body Text",
    )
    set_para_text(
        paras[204],
        "为判断 Actor 是否只是复制 prior，评估阶段额外记录 Actor 输出与 Lyapunov prior 融合后 portfolio 的差异 D_KL(w_i||p_i^L)。完整 FACR 的平均 D_KL 为 0.0234，高于 Prior-only 的 0.0104，说明 Actor/fusion 在 prior 附近学习了小幅 residual correction；同时 Prior-only 的 migration 约为完整 FACR 的 3.45 倍，说明学习型 proposal 与时间融合主要改善动作稳定性。",
        "Body Text",
    )
    for idx in [205, 206, 207, 208, 209]:
        remove_element(paras[idx]._p)
    remove_element(tables[12]._element)
    set_para_text(
        paras[210],
        "综合上述结果，FACR-MAPPO 更准确的定位是 stale Lyapunov-prior guided residual MAPPO：Lyapunov prior 提供主要队列风险控制，Actor 提供 residual correction，AgeGate 降低 feedback-age sensitivity，Q-penalty、concentration regularizer 和 max-step projection 则作为辅助正则或部署保护。该定位避免把每个模块都解释为独立主贡献，而强调完整方法在尾部风险、反馈年龄鲁棒性和动作稳定性之间取得更稳健的折中。",
        "Body Text",
    )

    # Robustness and efficiency merged into one section.
    set_para_text(paras[211], "5.5 跨环境鲁棒性与推理效率", "Heading 3")
    set_para_text(
        paras[212],
        "本节采用固定 checkpoint 鲁棒性协议：不在目标数据集或目标负载上重新训练，只评估已训练策略的泛化。该协议与第 5.3 节主 baseline 协议不同，因此绝对数值不能与 25-run 主对比直接横向比较。为避免混淆，正式跨环境汇总只使用 N∈{15,20,25,30}，KAIST N=40 单独作为压力测试报告。",
        "Body Text",
    )
    remove_element(paras[214]._p)
    remove_element(paras[215]._p)
    remove_element(paras[216]._p)
    remove_element(paras[217]._p)
    set_para_text(
        paras[218],
        "stale degradation 汇总图比较 δ=8 相对 δ=0 的性能退化。FACR 的相对退化最低，说明固定 checkpoint 策略对反馈陈旧度更稳健。",
        "Body Text",
    )
    remove_element(paras[222]._p)
    remove_element(paras[223]._p)
    remove_element(paras[224]._p)
    remove_element(tables[15]._element)
    set_para_text(
        paras[225],
        "推理效率方面，benchmark 使用 KAIST、Uniform、N=30、δ=8 和 10k 次测量。FACR 的 actor forward、prior/fusion/projection 总决策耗时为毫秒级，明显低于一个 1 s slot；QECO-A 由于每用户 DQN/LSTM 推理，平均决策时间更高但仍处于可部署范围。",
        "Body Text",
    )
    remove_element(paras[226]._p)

    # Discussion and conclusion renumbering, remove appendix note.
    set_para_text(paras[227], "6 讨论", "Heading 2")
    set_para_text(
        paras[229],
        "本文不宣称 FACR 具有全局稳定性保证。第 3.5 节的 mechanism analysis 建立在对称或近似同质服务器、局部平衡点附近、固定反馈年龄、用户响应可线性化、忽略 projection 和非线性策略细节等假设之上。它只能说明在局部模型中，反馈年龄和聚合用户响应强度增大会缩小系统稳定区域；实践含义是 stale feedback 在临界高利用率区比轻载区更危险，而不是证明完整神经策略全局稳定。",
        "Body Text",
    )
    set_para_text(paras[234], "7 结论", "Heading 2")
    remove_element(paras[236]._p)
    remove_element(paras[237]._p)

    # Renumber visible figure references/captions after deletion.
    renumber_figures(doc)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
