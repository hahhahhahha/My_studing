from __future__ import annotations

from pathlib import Path
from textwrap import dedent

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "results_priceedge_qtbias_paper_ready"
FIG = OUT / "figures"
MD_OUT = ROOT / "paper_draft_priceedge_qtbias_cn_v3.md"
DOCX_OUT = ROOT / "paper_draft_priceedge_qtbias_cn_v3.docx"
ARCH_FIG = FIG / "fig_priceedge_qtbias_architecture.png"


def _read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(OUT / name)


def _fmt(x, nd=3):
    if pd.isna(x):
        return ""
    if abs(float(x)) < 0.01 and float(x) != 0:
        return f"{float(x):.4f}"
    return f"{float(x):.{nd}f}"


def _md_table(df: pd.DataFrame, cols: list[tuple[str, str]], nd=3, max_rows: int | None = None) -> str:
    if max_rows:
        df = df.head(max_rows)
    header = "| " + " | ".join(label for _, label in cols) + " |"
    sep = "| " + " | ".join("---" for _ in cols) + " |"
    rows = []
    for _, r in df.iterrows():
        cells = []
        for key, _ in cols:
            v = r[key]
            if isinstance(v, (float, int)) and key not in {"runs"}:
                cells.append(_fmt(v, nd))
            else:
                cells.append(str(v))
        rows.append("| " + " | ".join(cells) + " |")
    return "\n".join([header, sep, *rows])


def create_architecture_figure() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    w, h = 1800, 850
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype("arial.ttf", 36)
        font = ImageFont.truetype("arial.ttf", 24)
        font_small = ImageFont.truetype("arial.ttf", 20)
    except Exception:
        font_title = font = font_small = ImageFont.load_default()

    colors = {
        "blue": "#2b6cb0",
        "green": "#2f855a",
        "orange": "#dd6b20",
        "pink": "#b83280",
        "gray": "#4a5568",
        "lightblue": "#ebf8ff",
        "lightgreen": "#f0fff4",
        "lightorange": "#fffaf0",
        "lightpink": "#fff5f7",
    }

    def box(x0, y0, x1, y1, text, fill, outline, title=False):
        d.rounded_rectangle([x0, y0, x1, y1], radius=22, fill=fill, outline=outline, width=4)
        lines = text.split("\n")
        f = font if title else font_small
        total_h = len(lines) * 28
        y = y0 + (y1 - y0 - total_h) / 2
        for line in lines:
            bbox = d.textbbox((0, 0), line, font=f)
            d.text((x0 + (x1 - x0 - (bbox[2] - bbox[0])) / 2, y), line, font=f, fill="#111111")
            y += 30

    def arrow(x0, y0, x1, y1, color="#333333", width=4):
        d.line([x0, y0, x1, y1], fill=color, width=width)
        # arrow head
        import math
        ang = math.atan2(y1 - y0, x1 - x0)
        L = 18
        for a in (ang + 2.6, ang - 2.6):
            d.line([x1, y1, x1 + L * math.cos(a), y1 + L * math.sin(a)], fill=color, width=width)

    headers = [
        (70, "1  Stale MEC observation", colors["blue"]),
        (520, "2  QT risk price", colors["orange"]),
        (965, "3  Edge-wise shared actor", colors["green"]),
        (1375, "4  Cooperative execution", colors["pink"]),
    ]
    for x, text, c in headers:
        d.rounded_rectangle([x, 35, x + 360, 88], radius=12, fill="#f7fafc", outline="#edf2f7")
        d.text((x + 18, 47), text, font=font, fill=c)

    box(80, 145, 410, 295, "Per-user state\nlocation, request load\nedge distances", colors["lightblue"], colors["blue"])
    box(80, 370, 410, 545, "Shared stale summary\nB(t-delta)\nwait, score, capacity\nfeedback age delta", colors["lightblue"], colors["blue"])
    box(535, 160, 850, 310, "QT price\npi_ij^QT = pi_ij^Q + pi_ij^T\nqueue pressure + delay", colors["lightorange"], colors["orange"])
    box(535, 390, 850, 545, "Price normalization\nrow min-max price\nprice scale s_i^pi", colors["lightorange"], colors["orange"])
    box(980, 130, 1290, 270, "User encoder\nh_i^u = f_u(user, age,\nprice scale)", colors["lightgreen"], colors["green"])
    box(980, 330, 1290, 500, "Shared edge encoder\nh_ij^e = f_e(d_ij, B_j,\npi_ij^QT, delta)", colors["lightgreen"], colors["green"])
    box(980, 565, 1290, 710, "Price-biased Gaussian\nmu_ij = m_ij^theta - lambda_pi pi_ij^QT\nw_i = softmax(z_i)", colors["lightgreen"], colors["green"])
    box(1390, 200, 1700, 360, "Portfolio action\nw_i(t) in simplex\nprobabilistic offloading", colors["lightpink"], colors["pink"])
    box(1390, 470, 1700, 650, "Joint environment\nA_j(t)=sum_i a_i w_ij\nqueue update Q_j(t+1)\nteam reward + MAPPO", colors["lightpink"], colors["pink"])

    arrow(410, 220, 535, 225, colors["blue"])
    arrow(410, 460, 535, 235, colors["blue"])
    arrow(700, 310, 700, 390, colors["orange"])
    arrow(850, 235, 980, 405, colors["orange"])
    arrow(850, 470, 980, 405, colors["orange"])
    arrow(1135, 270, 1135, 330, colors["green"])
    arrow(1135, 500, 1135, 565, colors["green"])
    arrow(1290, 640, 1390, 280, colors["green"])
    arrow(1545, 360, 1545, 470, colors["pink"])
    arrow(1390, 585, 1235, 710, "#e53e3e")
    d.text((1095, 745), "training / centralized critic / PPO update", font=font_small, fill="#e53e3e")

    d.text((80, 760), "Solid arrows: inference/action flow; red arrow: feedback/training flow.", font=font_small, fill=colors["gray"])
    img.save(ARCH_FIG)


def build_markdown() -> str:
    main = _read_csv("main_baseline_summary_no_facr.csv")
    ab = _read_csv("paper_ablation_summary_no_priceonly_localcritic.csv")
    paired = _read_csv("paper_ablation_paired_vs_main.csv")
    censor = _read_csv("qtbias_latency_cvar_censoring_audit_summary.csv")
    robust = _read_csv("fixed_checkpoint_robustness_summary_with_qtbias_no_facr.csv")
    n40 = _read_csv("n40_stress_summary_with_qtbias_no_facr.csv")

    main_table = _md_table(
        main,
        [
            ("method", "方法"),
            ("runs", "runs"),
            ("average_latency_mean", "Avg latency"),
            ("cvar99_mean", "CVaR99"),
            ("queue_oscillation_mean", "Queue osc."),
            ("portfolio_migration_rate_mean", "Migration"),
        ],
    )
    ab_table = _md_table(
        ab,
        [
            ("Variant", "变体"),
            ("runs", "runs"),
            ("average_latency_mean", "Avg latency"),
            ("cvar99_mean", "CVaR99"),
            ("queue_oscillation_mean", "Queue osc."),
            ("portfolio_migration_rate_mean", "Migration"),
            ("lcc_price_action_spearman_mean", "Spearman(-price,w)"),
        ],
    )
    paired_table = _md_table(
        paired,
        [
            ("Variant", "变体"),
            ("metric", "指标"),
            ("mean_diff", "paired diff"),
            ("ci95_low", "CI low"),
            ("ci95_high", "CI high"),
            ("relative_pct", "相对变化"),
            ("variant_better_rate_lower", "better rate"),
        ],
        nd=4,
        max_rows=10,
    )
    robust_table = _md_table(
        robust,
        [
            ("method", "方法"),
            ("runs", "runs"),
            ("average_latency_mean", "Avg latency"),
            ("cvar99_mean", "CVaR99"),
            ("queue_oscillation_mean", "Queue osc."),
            ("portfolio_migration_rate_mean", "Migration"),
        ],
    )
    n40_table = _md_table(
        n40,
        [
            ("method", "方法"),
            ("distribution", "分布"),
            ("runs", "runs"),
            ("average_latency_mean", "Avg latency"),
            ("cvar99_mean", "CVaR99"),
            ("queue_oscillation_mean", "Queue osc."),
            ("portfolio_migration_rate_mean", "Migration"),
        ],
    )
    c = censor.iloc[0]

    text = dedent(
        """
        # PriceEdge-QT-Bias MAPPO：共享滞后反馈下的多用户协作式 portfolio 卸载

        ## 摘要

        高并发移动边缘计算系统中，多个移动用户通常基于同一份周期广播的滞后服务器摘要进行卸载决策；当这些用户同步响应过期负载信息时，局部低负载指示可能被放大为队列尖峰、尾部时延和动作过冲。现有风险感知卸载、队列控制和不完整信息学习方法分别优化 tail risk、队列压力或探索利用，但较少把共享滞后反馈、多用户动作耦合与 portfolio 卸载统一为一个协作式多智能体控制问题。为此，本文提出 PriceEdge-QT-Bias MAPPO，将滞后服务器摘要转化为用户—服务器逐边队列—时延价格，并把该价格作为 action-aligned inductive bias 嵌入共享 edge-wise Gaussian Actor 的动作 logit，而不是用解析规则覆盖 MAPPO 策略输出。在 KAIST 主场景的 5 个训练种子 × 5 个 feedback age 网格上，PriceEdge-QT-Bias 取得 0.901 s 平均时延和 3.311 s CVaR99，较最强 adapted baseline PPO-MV-A 分别降低 29.3% 和 57.5%，同时保持较低队列振荡；固定 checkpoint 在 KAIST、NCSU、NewYork、Uniform/Gaussian 请求和不同用户规模下也保持最低总体平均时延与尾部风险。结果表明，共享滞后反馈下的 MEC 卸载应被视为用户动作耦合、状态陈旧度和队列演化耦合的协作控制问题；逐边风险价格与 MAPPO actor 的结构化耦合能够为 stale-feedback portfolio offloading 提供有效的可解释偏置。

        关键词：移动边缘计算；滞后状态反馈；多智能体强化学习；portfolio 卸载；尾部时延；风险价格

        ## 1 引言

        移动边缘计算（MEC）通过把计算和服务能力部署到靠近终端的位置，降低核心网传输压力并支撑实时移动服务 [1]–[7]。在移动感知、AR/VR、群智感知和实时智能服务中，用户任务需要在无线通信、边缘计算能力和服务器排队之间持续权衡。一个容易被简化但实际非常关键的问题是：用户在时隙 t 决策时通常并不掌握当前真实队列 Q(t)，而只能读取由边缘平台周期广播、同步或压缩后的服务器摘要 B(t-delta)。

        经典 stale load information 研究指出，旧负载信息不能被简单解释为当前真实负载；当多个请求者同时选择表面低负载节点时，过期信息可能导致负载均衡恶化 [24], [25]。Fog/edge 场景中的在线负载均衡研究进一步表明，网络传播延迟和状态相关性下降会恶化响应时间与负载分布 [26], [27]。因此，本文面对的不是一般“不完整信息”，而是一个更具体的控制难题：多个用户基于同一份滞后服务器摘要同时执行 portfolio action，可能诱发同步追逐、队列尖峰和系统性拥挤。

        本文将研究边界明确限定为：共享滞后反馈下的多用户协作式 portfolio 卸载。也就是说，本文不试图证明学习控制器优于所有可能的解析控制器；引入 MAPPO 的依据在于系统本身就是多智能体耦合决策过程。用户 i 输出 portfolio w_i(t)，而服务器 j 的聚合到达负载由所有用户共同决定：

        ```text
        A_j(t)=sum_{i=1}^N a_i(t) w_ij(t).
        ```

        因此，一个用户提高某台服务器的分配概率，会改变该服务器的聚合到达、后续队列、其他用户等待时间、系统尾部时延以及下一时隙所有用户所面对的状态。对单个用户而言，环境转移不是只由自身动作决定：

        ```text
        P_i(s_{t+1}|s_t,w_i) != P_i(s_{t+1}|s_t,w_i,W_-i).
        ```

        这就是本文采用协作式多智能体强化学习的直接动机。Price 在这一叙事中不是独立控制器，而是提供给多智能体策略的逐动作结构化风险先验，或者说 action-aligned inductive bias。本文核心问题可概括为：在多个移动用户同时基于同一份滞后服务器摘要进行概率化卸载时，如何通过协作式多智能体学习抑制同步追逐和队列尖峰，并联合优化平均时延、尾部时延与动作稳定性？

        本文贡献如下。

        1. 多智能体 stale-feedback 建模：将共享滞后服务器摘要、多用户 portfolio action 和服务器队列演化统一为协作式部分可观测随机博弈，突出用户动作耦合、共享滞后信息耦合和队列演化耦合。
        2. 逐边 QT 风险价格：从滞后等待压力和用户—服务器预测时延构造 pi_ij^QT=pi_ij^Q+pi_ij^T，把服务器级风险转化为动作维度的可解释方向。
        3. Price-aligned edge-wise Actor：建立 pi_ij -> mu_ij -> w_ij 的直接结构通道，使价格既作为逐服务器特征进入共享 edge encoder，又以负 logit bias 影响 Gaussian mean。
        4. 协作式风险训练与系统评估：使用共享 Actor、集中式 critic 和团队风险目标训练联合策略，并在多 seed、多 feedback age、跨数据集、跨请求分布和高负载压力下评估平均时延、CVaR99、队列振荡与动作迁移。

        ## 2 相关工作

        ### 2.1 MEC 卸载与高并发边缘服务

        MEC 综述和标准化工作强调，边缘侧计算、缓存和服务编排能够降低端到端时延 [1]–[7]。多数卸载模型把动作定义为本地执行、单服务器选择或连续资源分配。本文使用 portfolio action：每个用户输出 w_i(t) in Delta^(M-1)，表示其任务在多个边缘服务器间的概率化或加权分配。这种动作空间更适合高并发场景，但也使多用户相关动作和系统性拥挤更突出。

        ### 2.2 Stale load information 与 delayed feedback

        Dahlin [24] 和 Mitzenmacher [25] 对旧负载信息的有效性和误导性进行了经典分析，指出旧状态不能被当作当前负载真值。Beraldi 等在 fog/edge computing 中研究 loosely correlated state information 和 stale information 对在线负载均衡的影响 [26], [27]。Tahir 等把延迟反馈下的计算集群负载均衡建模为部分可观测控制问题 [28]。这些工作共同支撑本文的核心假设：服务器摘要 B(t-delta) 仍有统计参考价值，但必须被视为低置信度、滞后的结构信息。

        ### 2.3 尾部时延、风险敏感卸载与队列控制

        Tail latency 文献说明，在大规模服务系统中，少量慢请求会显著影响整体体验 [8], [9]。风险敏感边缘卸载工作使用 DSR、mean-variance 或 CVaR 控制低概率高时延事件 [18], [30], [31]。另一方面，Lyapunov drift-plus-penalty 和 backpressure/max-weight 是随机网络队列稳定控制的经典工具 [10], [11]，并已被用于 MEC 动态卸载和能量—时延权衡 [12], [13]。本文继承这些思想，但不把 Lyapunov 规则作为执行时的外部 action override，而是把队列—时延信息转化为 MAPPO actor 内部的动作对齐偏置。

        ### 2.4 信息新鲜度与服务器摘要年龄

        AoI/AoT 相关工作说明，信息年龄本身会影响 MEC 决策 [14]–[17], [29]。但这类研究通常关注任务数据、监测结果或计算结果的新鲜度，而本文关注的是用户做卸载决策时使用的服务器负载摘要是否新鲜。换言之，本文研究的是 decision information freshness，而非 task-result freshness。

        ### 2.5 Baseline 适配边界

        本文对比的 QECO-A、LyFOTO-LC、TODA-MAB-A、PPO-DSR-A 和 PPO-MV-A 均是 adapted baseline：保留原方法的核心思想，但迁移到统一 portfolio action、stale feedback、训练预算和评估环境中。它们不是原论文完整系统逐项复刻。这样做的目的是在统一问题接口下检验已有思想在 stale portfolio offloading 问题中的可迁移性；本文不会将这些结果解释为对原论文完整系统的否定。

        ## 3 系统模型与问题定义

        ### 3.1 协作式部分可观测随机博弈

        系统包含 N 个移动用户和 M 个边缘服务器。时间被划分为长度为 Delta t 的离散时隙。用户 i 在时隙 t 的位置为 x_i(t)，边缘服务器 j 的位置和服务率分别为 y_j 与 mu_j(t)。用户只能观察局部状态和滞后广播摘要：

        ```text
        o_i(t)=[x_i(t), r_i(t), d_i(t), B(t-delta), delta].
        ```

        其中 d_i(t) 是用户到各边缘服务器的距离向量，B(t-delta) 包含服务器评分、等待时间、队列摘要、容量和反馈年龄。本文将问题表述为协作式部分可观测随机博弈：

        ```text
        G=<N,S,{O_i},{A_i},P,R,gamma>.
        ```

        每个用户是一个 agent，动作空间为概率单纯形 Delta^(M-1)，奖励为团队共享的风险感知目标。Actor 执行时不访问当前真实队列 Q(t)；critic 在训练时可以使用全局状态以降低多智能体非平稳性。

        ### 3.2 请求、portfolio action 与任务分配

        每个用户在每个时隙产生随机任务数：

        ```text
        K_i(t) ~ D_A.
        ```

        用户 i 的 portfolio 为：

        ```text
        w_i(t)=[w_i1(t),...,w_iM(t)] in Delta^(M-1).
        ```

        当前实现把 portfolio 解释为概率化卸载：用户的每个任务按照 Categorical(w_i(t)) 抽样服务器：

        ```text
        J_ik(t) ~ Categorical(w_i(t)).
        ```

        若任务 k 的计算工作量为 c_ik，则服务器 j 的聚合到达负载为：

        ```text
        A_j(t)=sum_i sum_{k=1}^{K_i(t)} c_ik * 1{{J_ik(t)=j}}.
        ```

        在期望意义下，这等价于 A_j(t)=sum_i a_i(t)w_ij(t)，其中 a_i(t)=sum_k c_ik。该定义说明，单个用户动作的后果由所有用户的联合 portfolio 共同决定。

        ### 3.3 时延、队列更新与边界处理

        用户 i 的任务若被分配到服务器 j，其时延由通信、排队和服务三部分组成：

        ```text
        T_ij(t)=T_ij^comm(t)+T_j^queue(t)+T_j^service(t).
        ```

        当前仿真中，通信时延随用户—服务器距离线性增长；服务时延由任务计算量和服务器服务率决定；排队时延由服务器等待时间或队列长度换算得到。服务器队列更新为：

        ```text
        Q_j(t+1)=min{{Q_j^max, [Q_j(t)-mu_j(t)Delta t]^+ + A_j(t)}}.
        ```

        队列溢出采用系统上限裁剪；若任务时延超过仿真上限或产生 timeout，需在结果审计中同时报告 capped/uncapped latency、timeout 数量、deadline violation probability 和 CVaR 尾部样本数。本文主方法的主场景审计显示未出现 T>=10s 的右删失，因此 CVaR99 不是由 10 秒上限饱和造成。

        ### 3.4 评价指标

        Average latency 为测试阶段所有完成任务的平均端到端时延。CVaR99 定义为任务级 latency 分布最差 1% 样本的均值。Queue oscillation 衡量服务器队列压力随时间的波动。Global concentration 使用归一化 HHI 描述系统负载是否过度集中。Migration 表示相邻时隙 portfolio 的平均 L1 变化：

        ```text
        Mig = mean_{i,t} ||w_i(t)-w_i(t-1)||_1.
        ```

        因此 migration 衡量的是策略输出的动作迁移/切换成本，而不是已执行任务在服务器间的物理迁移。

        ## 4 PriceEdge-QT-Bias MAPPO 方法

        ![Fig. 1. PriceEdge-QT-Bias MAPPO algorithm architecture](@@ARCH_FIG@@)

        图 1 展示了本文方法结构。算法主线为：

        ```text
        B(t-delta) -> pi_ij^QT -> edge-wise shared actor -> w_i(t) -> A_j(t) -> Q_j(t+1).
        ```

        ### 4.1 逐边队列—时延价格

        PriceEdge-QT-Bias 使用滞后服务器摘要构造两个动作维度价格。首先，队列压力价格为：

        ```text
        pi_ij^Q(t)=a_bar_i(t) W_j^delta,norm(t),
        ```

        其中 a_bar_i(t) 是归一化用户工作量，W_j^delta,norm(t) 是服务器 j 的滞后归一化等待压力。该项表达“大任务用户继续选择拥塞服务器”的风险。

        第二，基础时延价格为：

        ```text
        pi_ij^T(t)=V tau_ij(t),
        ```

        其中 tau_ij(t) 是由当前距离、服务能力和滞后等待时间构成的归一化预测时延，V 是性能权重。最终 QT 价格为：

        ```text
        pi_ij^QT(t)=pi_ij^Q(t)+pi_ij^T(t).
        ```

        该价格不是执行时的解析控制器，也不直接决定 portfolio；它提供局部拥塞与时延方向，供 MAPPO 策略学习在长期联合演化中如何使用。

        ### 4.2 价格归一化与价格尺度

        为避免不同用户任务量和距离尺度导致价格范围不可比，代码对每个用户的价格向量执行逐行 min-max normalization：

        ```text
        pi_ij <- (pi_ij - min_k pi_ik)/(max_k pi_ik - min_k pi_ik + eps).
        ```

        同时记录绝对价格尺度：

        ```text
        s_i^pi = log(1 + mean_j [pi_ij^raw]^+).
        ```

        因此，pi_i 告诉 actor “应该避开哪台服务器”，s_i^pi 则告诉 actor “当前风险压力有多强”。

        ### 4.3 逐服务器共享 Actor 与价格 logit bias

        普通 flat actor 通常把所有服务器特征拼接成长向量，再输出 M 维动作。该结构需要网络自行学习“第 j 个服务器特征对应第 j 个动作 logit”。PriceEdge-QT-Bias 改用逐服务器共享 actor。用户级编码为：

        ```text
        h_i^u=f_u(user_i, feedback_age, s_i^pi).
        ```

        对每个服务器 j，构造逐边特征：

        ```text
        e_ij(t)=[d_ij(t), B_j(t-delta), pi_ij^QT(t), delta, s_i^pi(t)].
        ```

        共享服务器编码器和打分器为：

        ```text
        h_ij^e=f_e(e_ij),  m_ij^theta=f_s(h_i^u,h_ij^e).
        ```

        价格以负偏置形式进入 Gaussian mean：

        ```text
        mu_ij(t)=m_ij^theta(t)-lambda_pi pi_ij^QT(t).
        z_i(t) ~ N(mu_i(t), sigma_i^2(t)),  w_i(t)=softmax(z_i(t)).
        ```

        代码中 price 同时作为逐服务器特征进入 edge encoder，并以固定负偏置作用于 Gaussian mean。因此严格地说，price-to-logit bias 保证了相对于 lambda_pi 的方向性，而不是保证总导数 d mu_ij / d pi_ij 恒为负；本文据此命名为 PriceEdge-QT-Bias MAPPO，而不称为严格单调控制器。

        ### 4.4 协作式训练目标

        所有用户共享 actor 参数。训练采用 CTDE：actor 执行时只使用局部观测和滞后摘要，centralized critic 在训练时读取系统级状态、队列和联合动作统计。团队奖励综合平均时延、尾部风险、队列振荡和动作迁移。本文把 Lyapunov/CVaR 思想作为 bounded team reward 和风险状态的训练信号，而不是声称实现了严格漂移上界最小化或全局队列稳定性证明。

        在线执行流程如下。

        ```text
        Algorithm 1: Online PriceEdge-QT-Bias execution
        Input: user state x_i(t), request load a_i(t), stale summary B(t-delta), age delta
        1. For each edge j, compute predicted delay tau_ij(t).
        2. Compute QT risk price pi_ij^QT(t)=pi_ij^Q(t)+pi_ij^T(t).
        3. Normalize price vector pi_i and compute price scale s_i^pi.
        4. Use shared edge-wise actor to compute m_ij^theta.
        5. Apply price bias mu_ij=m_ij^theta-lambda_pi pi_ij^QT.
        6. Sample Gaussian latent z_i and output w_i(t)=softmax(z_i).
        7. Execute probabilistic portfolio offloading.
        ```

        ```text
        Algorithm 2: CTDE training
        1. Collect rollout tuples from all users under shared actor.
        2. Execute joint portfolio actions and update queues Q_j(t+1).
        3. Compute team reward from latency, tail risk, queue pressure and migration.
        4. Estimate advantages with centralized critic.
        5. Update shared actor and critic using clipped PPO objective.
        ```

        ## 5 实验设置

        主实验使用 KAIST 轨迹，N=30 个用户，M=35 个候选边缘服务器，Poisson-bursty 请求，feedback age delta in {{0,1,2,4,8}}。每个方法使用 5 个训练种子，每个 seed-delay 单元训练 10k steps 并测试 1000 steps。正式 baseline 均使用相同动作空间、反馈接口和评估指标。PriceEdge-QT-Bias 的主要参数为：V=0.5，lambda_m=0.05，lambda_pi=0.75，target portfolio entropy=0.99，QT price components={{queue, latency}}。

        ## 6 实验结果

        ### 6.1 主 baseline 对比

        表 1 给出 KAIST 主协议下的正式结果。PriceEdge-QT-Bias 在 average latency、CVaR99 和 queue oscillation 上均优于 adapted baselines；与最强平均时延 baseline PPO-MV-A 相比，平均时延降低约 29.3%，CVaR99 降低约 57.5%。Migration 不是最低，说明该方法以一定 portfolio 调整换取更低时延和尾部风险。

        表 1. 主 baseline 多指标对比（5 seeds × 5 feedback ages）。

        @@MAIN_TABLE@@

        ![Fig. 2. Main baseline comparison](@@FIG_MAIN@@)

        ### 6.2 Feedback age 分解

        图 3 按 feedback age 分解平均时延和 CVaR99。该实验用于验证共享滞后摘要是否会随 delta 增大而改变策略表现，也用于检查 PriceEdge-QT-Bias 在 fresh、中等滞后和严重滞后反馈下是否保持稳定。由于每个 delay 是固定协议下独立训练，本文不把非单调曲线解释为单一策略的在线 age 自适应，而只把它作为不同滞后条件下的公平性能剖面。

        ![Fig. 3. Performance by feedback delay](@@FIG_DELAY@@)

        ### 6.3 结构消融：价格、对齐和逐边结构

        主消融围绕算法内部机制展开：显式 price-to-logit bias 是否必要、price 与服务器动作是否必须正确对齐、edge-wise actor 是否优于 flat actor、stale margin 和 tail state 是否构成核心贡献。

        表 2. PriceEdge-QT-Bias 结构消融。

        @@AB_TABLE@@

        ![Fig. 4. Ablation comparison](@@FIG_ABLATION@@)

        Paired bootstrap 结果显示，QT-NoBias 相比完整方法使 CVaR99 增加约 3.06%，QT-Bias-Shuffle 使 CVaR99 增加约 4.58%，说明显式负价格偏置和正确 price-server 对齐是有贡献的。Flat-Bias-QT 的平均时延和 CVaR99 均弱于完整方法，支持逐服务器共享 actor 的结构价值。另一方面，w/o stale margin 与完整方法几乎一致，w/o tail state Z 甚至略好，说明这两个模块不应被写作核心贡献；当前 paper-facing 方法应强调 QT 价格、logit bias、edge-wise actor 和协作式 MAPPO 训练。

        表 3. 相对完整方法的 paired difference（正值表示变体更差）。

        @@PAIRED_TABLE@@

        ### 6.4 同 checkpoint 的价格干预

        为排除“不同训练随机性”对消融解释的干扰，本文还使用同一个训练完成的 PriceEdge-QT-Bias checkpoint，在测试时进行价格干预：aligned price、shuffled price、zero price 和 constant price。该诊断不改变网络权重，只改变测试时价格输入/偏置，因此更直接地检验策略是否依赖正确的逐边价格信息。

        ![Fig. 5. Test-time price intervention](@@FIG_INTERVENTION@@)

        ### 6.5 尾部分布与 CVaR 截断审计

        图 6 给出代表性 severe stale 场景下的 latency CCDF。除均值和 CVaR99 外，CCDF 能显示右尾概率是否真正下降，避免只用单个统计量掩盖分布形态。

        ![Fig. 6. Latency CCDF under severe stale feedback](@@FIG_CCDF@@)

        为检查 CVaR99 是否受到 10 秒上限、timeout penalty 或右删失影响，本文读取 PriceEdge-QT-Bias 测试阶段逐任务 latency。主场景平均每个 seed-delay 单元包含 @@TASKS_MEAN@@ 个任务，最差 1% 平均包含 @@TAIL99_MEAN@@ 个任务；uncapped CVaR99 为 @@CVAR99_UNCAPPED@@ s，capped@10s CVaR99 为 @@CVAR99_CAPPED@@ s，P(T>=10s) 为 @@PROB_GE10@@%。因此，PriceEdge-QT-Bias 的主结果不是由 10 秒右删失或 timeout 饱和造成。

        ### 6.6 固定 checkpoint 跨环境测试

        为避免把跨环境测试变成重新调参，本文使用固定 KAIST 主场景 checkpoint，在 KAIST、NCSU、NewYork，Uniform/Gaussian 请求，N in {{15,20,25,30}}，delta in {{0,8}} 上直接评估。该协议常用于检验已训练策略的外推和鲁棒性，但不等价于“在每个新环境重新训练后的最优性能”。

        表 4. 固定 checkpoint 跨环境测试汇总。

        @@ROBUST_TABLE@@

        ![Fig. 7. Fixed-checkpoint robustness](@@FIG_ROBUST@@)

        ### 6.7 KAIST N=40 高负载压力测试

        N=40 用作额外 high-load stress test，而不是主协议。表 5 和图 8 显示，当系统接近容量边界时，各方法时延和 CVaR99 均明显升高；PriceEdge-QT-Bias 仍能避免 PPO-DSR-A/PPO-MV-A 接近 10s CVaR 饱和的情况，但该结果应解释为相对抗压能力，而非持续过载下的绝对服务保证。

        表 5. KAIST N=40 pressure test。

        @@N40_TABLE@@

        ![Fig. 8. KAIST N=40 stress test](@@FIG_N40@@)

        ## 7 讨论

        当前证据支持三个相对稳健的结论。第一，共享滞后反馈下的 portfolio 卸载确实是多用户耦合问题，MAPPO 的作用是学习联合策略，而不是简单修正某个解析控制器。第二，QT price-to-logit bias 与正确 price-server 对齐在多 seed 消融中均能降低 CVaR99，说明 price 作为 action-aligned inductive bias 是有价值的。第三，edge-wise actor 比 flat actor 更适合该任务，因为它显式绑定了服务器特征、价格和动作 logit。

        同时，本文不夸大当前版本的机制。w/o stale margin 与 w/o tail state Z 没有变差，说明服务器特定 stale uncertainty margin 和 tail-risk state 目前不是核心增益来源。因此，本文仅将 CTDE 作为训练实现，而不把 critic 信息结构列为主要创新。未来更值得补充的是非统一 feedback age、异步用户更新、missed broadcast、capacity shock 和 hidden background load，以检验逐边结构在更复杂信息不一致条件下的价值。

        ## 8 结论

        本文研究共享滞后反馈下的多用户协作式 portfolio 卸载，并提出 PriceEdge-QT-Bias MAPPO。该方法把滞后队列压力和用户—服务器预测时延转化为逐边 QT 风险价格，并通过 edge-wise shared actor 和显式 logit bias 将价格嵌入 MAPPO 策略内部。实验表明，PriceEdge-QT-Bias 在统一 stale-summary portfolio offloading 接口下显著降低 adapted baselines 的平均时延和 CVaR99，并在固定 checkpoint 跨环境测试与 N=40 压力测试中保持相对优势。未来工作将面向异步 feedback age、隐藏背景负载和真实 MEC 状态日志，进一步验证逐边风险价格在非平稳边缘系统中的泛化能力。

        ## 参考文献

        [1] ETSI, “Mobile Edge Computing—Introductory Technical White Paper,” 2014.  
        [2] W. Shi et al., “Edge Computing: Vision and Challenges,” IEEE Internet of Things Journal, 2016.  
        [3] M. Satyanarayanan, “The Emergence of Edge Computing,” Computer, 2017.  
        [4] N. Abbas et al., “Mobile Edge Computing: A Survey,” IEEE Internet of Things Journal, 2018.  
        [5] P. Mach and Z. Becvar, “Mobile Edge Computing: A Survey on Architecture and Computation Offloading,” IEEE Communications Surveys & Tutorials, 2017.  
        [6] Y. Mao et al., “A Survey on Mobile Edge Computing: The Communication Perspective,” IEEE Communications Surveys & Tutorials, 2017.  
        [7] L. Lin et al., “Computation Offloading Toward Edge Computing,” Proceedings of the IEEE, 2019.  
        [8] J. Dean and L. A. Barroso, “The Tail at Scale,” Communications of the ACM, 2013.  
        [9] A. Sriraman et al., “Deconstructing the Tail at Scale Effect Across Network Protocols,” arXiv:1701.03100, 2017.  
        [10] M. J. Neely, Stochastic Network Optimization with Application to Communication and Queueing Systems, 2010.  
        [11] L. Tassiulas and A. Ephremides, “Stability Properties of Constrained Queueing Systems and Scheduling Policies,” IEEE Transactions on Automatic Control, 1992.  
        [12] Y. Mao et al., “Dynamic Computation Offloading for Mobile-Edge Computing with Energy Harvesting Devices,” IEEE JSAC, 2016.  
        [13] Y. Mao et al., “Power-Delay Tradeoff in Multi-User Mobile-Edge Computing Systems,” IEEE JSAC, 2016.  
        [14] X. Song et al., “Age Based Task Scheduling and Computation Offloading in Mobile-Edge Computing Systems,” arXiv:1905.11570, 2019.  
        [15] Z. Tang et al., “Age of Information of Multi-user Mobile Edge Computing Systems,” arXiv:2211.00874, 2022.  
        [16] S. Chen et al., “Minimizing Age of Information for Mobile Edge Computing Systems,” arXiv:2307.01366, 2023.  
        [17] Y. Dong et al., “Mean Age of Information in Partial Offloading Mobile Edge Computing Networks,” arXiv:2409.16115, 2024.  
        [18] “Towards Risk-Averse Edge Computing With Deep Reinforcement Learning,” local baseline manuscript.  
        [19] J. Schulman et al., “Proximal Policy Optimization Algorithms,” arXiv:1707.06347, 2017.  
        [20] C. Yu et al., “The Surprising Effectiveness of PPO in Cooperative Multi-Agent Games,” NeurIPS, 2022.  
        [21] QECO, “QoE-oriented distributed DRL MEC offloading,” IEEE TNSE, 2025.  
        [22] LyFOTO, “Lyapunov + federated DRL + queue constraints for MEC,” IEEE TMC, 2025.  
        [23] TODAuction-MAB, “MEC incomplete information / MAB offloading,” Cluster Computing, 2025.  
        [24] M. Dahlin, “Interpreting Stale Load Information,” IEEE TPDS, 2000.  
        [25] M. Mitzenmacher, “How Useful Is Old Information?” IEEE TPDS, 2000.  
        [26] R. Beraldi et al., “Randomized Load Balancing under Loosely Correlated State Information in Fog Computing,” ACM MSWiM, 2020.  
        [27] R. Beraldi et al., “On the Impact of Stale Information on Distributed Online Load Balancing Protocols for Edge Computing,” Computer Networks, 2022.  
        [28] A. Tahir et al., “Load Balancing in Compute Clusters With Delayed Feedback,” arXiv:2207.03465, 2022.  
        [29] X. Chen et al., “Information Freshness-Aware Task Offloading in Air-Ground Integrated Edge Computing Systems,” IEEE JSAC, 2022.  
        [30] X. Hao et al., “A Risk-Sensitive Task Offloading Strategy for Edge Computing in Industrial Internet of Things,” EURASIP JWCN, 2021.  
        [31] “Robust Risk-Sensitive Task Offloading for Edge-Enabled Industrial Internet of Things,” 2024.
        """
    ).strip() + "\n"
    replacements = {
        "@@ARCH_FIG@@": ARCH_FIG.as_posix(),
        "@@FIG_MAIN@@": (FIG / "fig_qtbias_main_baseline_no_facr.png").as_posix(),
        "@@FIG_DELAY@@": (FIG / "fig_qtbias_main_by_delay_no_facr.png").as_posix(),
        "@@FIG_ABLATION@@": (FIG / "fig_qtbias_ablation_no_priceonly_localcritic.png").as_posix(),
        "@@FIG_INTERVENTION@@": (FIG / "fig_qtbias_testtime_price_intervention.png").as_posix(),
        "@@FIG_CCDF@@": (FIG / "fig_qtbias_latency_ccdf_delay8.png").as_posix(),
        "@@FIG_ROBUST@@": (FIG / "fig_qtbias_fixed_checkpoint_robustness.png").as_posix(),
        "@@FIG_N40@@": (FIG / "fig_qtbias_n40_stress.png").as_posix(),
        "@@MAIN_TABLE@@": main_table,
        "@@AB_TABLE@@": ab_table,
        "@@PAIRED_TABLE@@": paired_table,
        "@@ROBUST_TABLE@@": robust_table,
        "@@N40_TABLE@@": n40_table,
        "@@TASKS_MEAN@@": f"{c['n_test_tasks_mean']:.0f}",
        "@@TAIL99_MEAN@@": f"{c['tail99_sample_count_mean']:.0f}",
        "@@CVAR99_UNCAPPED@@": f"{c['cvar99_uncapped_mean']:.3f}",
        "@@CVAR99_CAPPED@@": f"{c['cvar99_capped_10s_mean']:.3f}",
        "@@PROB_GE10@@": f"{100*c['prob_latency_ge_10s_mean']:.4f}",
    }
    for key, value in replacements.items():
        text = text.replace(key, value)
    return text


def add_table(doc: Document, df: pd.DataFrame, cols: list[tuple[str, str]], nd=3, max_rows: int | None = None) -> None:
    if max_rows:
        df = df.head(max_rows)
    table = doc.add_table(rows=1, cols=len(cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, (_, label) in enumerate(cols):
        hdr[i].text = label
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, (key, _) in enumerate(cols):
            v = row[key]
            if isinstance(v, (float, int)) and key not in {"runs"}:
                cells[i].text = _fmt(v, nd)
            else:
                cells[i].text = str(v)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8.5)


def add_fig(doc: Document, path: Path, caption: str, width=6.2) -> None:
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True


def build_docx(markdown: str) -> None:
    main = _read_csv("main_baseline_summary_no_facr.csv")
    ab = _read_csv("paper_ablation_summary_no_priceonly_localcritic.csv")
    paired = _read_csv("paper_ablation_paired_vs_main.csv")
    censor = _read_csv("qtbias_latency_cvar_censoring_audit_summary.csv").iloc[0]
    robust = _read_csv("fixed_checkpoint_robustness_summary_with_qtbias_no_facr.csv")
    n40 = _read_csv("n40_stress_summary_with_qtbias_no_facr.csv")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Microsoft YaHei"

    def para(text: str, align=None):
        p = doc.add_paragraph(text)
        if align:
            p.alignment = align
        return p

    title = doc.add_heading("PriceEdge-QT-Bias MAPPO：共享滞后反馈下的多用户协作式 portfolio 卸载", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("摘要", 1)
    para("高并发移动边缘计算系统中，多个移动用户通常基于同一份周期广播的滞后服务器摘要进行卸载决策；当这些用户同步响应过期负载信息时，局部低负载指示可能被放大为队列尖峰、尾部时延和动作过冲。现有风险感知卸载、队列控制和不完整信息学习方法分别优化 tail risk、队列压力或探索利用，但较少把共享滞后反馈、多用户动作耦合与 portfolio 卸载统一为一个协作式多智能体控制问题。为此，本文提出 PriceEdge-QT-Bias MAPPO，将滞后服务器摘要转化为用户—服务器逐边队列—时延价格，并把该价格作为 action-aligned inductive bias 嵌入共享 edge-wise Gaussian Actor 的动作 logit，而不是用解析规则覆盖 MAPPO 策略输出。在 KAIST 主场景的 5 个训练种子 × 5 个 feedback age 网格上，PriceEdge-QT-Bias 取得 0.901 s 平均时延和 3.311 s CVaR99，较最强 adapted baseline PPO-MV-A 分别降低 29.3% 和 57.5%，同时保持较低队列振荡；固定 checkpoint 在 KAIST、NCSU、NewYork、Uniform/Gaussian 请求和不同用户规模下也保持最低总体平均时延与尾部风险。")
    para("关键词：移动边缘计算；滞后状态反馈；多智能体强化学习；portfolio 卸载；尾部时延；风险价格")

    sections = [
        ("1 引言", [
            "本文将研究边界限定为共享滞后反馈下的多用户协作式 portfolio 卸载，而不是学习控制器与所有可能解析控制器的比较。MAPPO 的建模动机来自多用户动作耦合：每个用户的 portfolio 会改变服务器聚合到达、队列演化和其他用户等待时间。",
            "Price 在本文中不是独立控制器，而是提供给多智能体策略的逐动作结构化风险先验。Price 负责局部动作方向，MAPPO 负责联合策略学习。"
        ]),
        ("2 相关工作", [
            "本文相关工作围绕 MEC 卸载、stale load information、tail latency/risk-aware offloading、Lyapunov 队列控制和信息新鲜度展开。Dahlin、Mitzenmacher 与 Beraldi 等 stale-load 文献直接支撑“旧负载摘要可能误导同步请求者”的问题动机；AoI/AoT 文献提供“信息年龄具有决策意义”的邻近证据；风险敏感卸载文献说明平均时延不能替代 CVaR99。"
        ]),
        ("3 系统模型与问题定义", [
            "系统建模为协作式部分可观测随机博弈 G=<N,S,{O_i},{A_i},P,R,gamma>。用户只观察 o_i(t)=[x_i(t), r_i(t), d_i(t), B(t-delta), delta]，Actor 执行时不访问当前真实队列 Q(t)。",
            "任务数 K_i(t)~D_A，任务服务器 J_ik(t)~Categorical(w_i(t))，聚合负载 A_j(t)=sum_i sum_k c_ik 1{J_ik(t)=j}。队列更新为 Q_j(t+1)=min{Q_j^max,[Q_j(t)-mu_j(t)Delta t]^+ + A_j(t)}。Migration 表示相邻时隙 portfolio 的 L1 变化，而不是真实任务迁移。"
        ]),
        ("4 PriceEdge-QT-Bias MAPPO 方法", [
            "算法链条为 B(t-delta) -> pi_ij^QT -> edge-wise shared actor -> w_i(t) -> A_j(t) -> Q_j(t+1)。QT 价格 pi_ij^QT=pi_ij^Q+pi_ij^T，其中 pi_ij^Q=a_bar_i(t)W_j^delta,norm(t)，pi_ij^T=V tau_ij(t)。",
            "逐服务器 Actor 使用共享编码：h_i^u=f_u(user_i,feedback age,s_i^pi)，h_ij^e=f_e(d_ij,B_j(t-delta),pi_ij^QT,delta,s_i^pi)，m_ij^theta=f_s(h_i^u,h_ij^e)。价格以负偏置进入 Gaussian mean：mu_ij=m_ij^theta-lambda_pi pi_ij^QT，随后 z_i~N(mu_i,sigma_i^2)，w_i=softmax(z_i)。"
        ]),
    ]
    for h, ps in sections:
        doc.add_heading(h, 1)
        for p in ps:
            para(p)
        if h.startswith("4 "):
            add_fig(doc, ARCH_FIG, "Fig. 1. PriceEdge-QT-Bias MAPPO algorithm architecture.", width=6.8)

    doc.add_heading("5 实验设置", 1)
    para("主实验使用 KAIST 轨迹，N=30，M=35，Poisson-bursty 请求，feedback age delta={0,1,2,4,8}。每个方法使用 5 个训练种子，每个 seed-delay 单元训练 10k steps 并测试 1000 steps。PriceEdge-QT-Bias 的主要参数为 V=0.5，lambda_m=0.05，lambda_pi=0.75，target portfolio entropy=0.99，price components={queue, latency}。")

    doc.add_heading("6 实验结果", 1)
    doc.add_heading("6.1 主 baseline 对比", 2)
    para("PriceEdge-QT-Bias 在 average latency、CVaR99 和 queue oscillation 上均优于 adapted baselines。与最强平均时延 baseline PPO-MV-A 相比，平均时延降低约 29.3%，CVaR99 降低约 57.5%。")
    para("表 1. 主 baseline 多指标对比（5 seeds × 5 feedback ages）。", WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, main, [("method", "方法"), ("runs", "runs"), ("average_latency_mean", "Avg latency"), ("cvar99_mean", "CVaR99"), ("queue_oscillation_mean", "Queue osc."), ("portfolio_migration_rate_mean", "Migration")])
    add_fig(doc, FIG / "fig_qtbias_main_baseline_no_facr.png", "Fig. 2. Main baseline comparison.")

    doc.add_heading("6.2 Feedback age 分解", 2)
    para("Feedback age 分解用于检查不同滞后条件下的公平性能剖面。由于每个 delay 是固定协议下独立训练，本文不把非单调曲线解释为单一策略的在线 age 自适应。")
    add_fig(doc, FIG / "fig_qtbias_main_by_delay_no_facr.png", "Fig. 3. Performance by feedback delay.")

    doc.add_heading("6.3 结构消融", 2)
    para("主消融聚焦显式 price-to-logit bias、price-server 对齐、edge-wise actor、stale margin 和 tail state。")
    para("表 2. PriceEdge-QT-Bias 结构消融。", WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, ab, [("Variant", "变体"), ("runs", "runs"), ("average_latency_mean", "Avg latency"), ("cvar99_mean", "CVaR99"), ("queue_oscillation_mean", "Queue osc."), ("portfolio_migration_rate_mean", "Migration"), ("lcc_price_action_spearman_mean", "Spearman")])
    add_fig(doc, FIG / "fig_qtbias_ablation_no_priceonly_localcritic.png", "Fig. 4. Ablation comparison.")
    para("Paired bootstrap 表明 QT-NoBias 与 QT-Bias-Shuffle 均恶化 CVaR99，支持显式价格偏置和正确 price-server 对齐；Flat-Bias-QT 弱于完整方法，支持逐服务器共享 actor。w/o stale margin 与 w/o tail state Z 没有变差，因此二者不应写作核心贡献。")
    para("表 3. 相对完整方法的 paired difference（正值表示变体更差）。", WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, paired.head(10), [("Variant", "变体"), ("metric", "指标"), ("mean_diff", "diff"), ("ci95_low", "CI low"), ("ci95_high", "CI high"), ("relative_pct", "相对变化"), ("variant_better_rate_lower", "better rate")], nd=4)

    doc.add_heading("6.4 同 checkpoint 价格干预", 2)
    para("同 checkpoint 干预固定网络权重，只改变测试时价格输入/偏置，用来检验策略是否依赖正确逐边价格信息。")
    add_fig(doc, FIG / "fig_qtbias_testtime_price_intervention.png", "Fig. 5. Test-time price intervention.")

    doc.add_heading("6.5 尾部分布与 CVaR 截断审计", 2)
    add_fig(doc, FIG / "fig_qtbias_latency_ccdf_delay8.png", "Fig. 6. Latency CCDF under severe stale feedback.")
    para(f"主场景平均每个 seed-delay 单元包含 {censor['n_test_tasks_mean']:.0f} 个任务，最差 1% 平均包含 {censor['tail99_sample_count_mean']:.0f} 个任务；uncapped CVaR99 为 {censor['cvar99_uncapped_mean']:.3f} s，capped@10s CVaR99 为 {censor['cvar99_capped_10s_mean']:.3f} s，P(T>=10s) 为 {100*censor['prob_latency_ge_10s_mean']:.4f}%。因此主结果不是由 10 秒右删失或 timeout 饱和造成。")

    doc.add_heading("6.6 固定 checkpoint 跨环境测试", 2)
    para("固定 KAIST 主场景 checkpoint，在 KAIST、NCSU、NewYork，Uniform/Gaussian 请求，N={15,20,25,30}，delta={0,8} 上直接评估。该协议用于检验外推能力，不等价于每个新环境重新训练后的最优性能。")
    para("表 4. 固定 checkpoint 跨环境测试汇总。", WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, robust, [("method", "方法"), ("runs", "runs"), ("average_latency_mean", "Avg latency"), ("cvar99_mean", "CVaR99"), ("queue_oscillation_mean", "Queue osc."), ("portfolio_migration_rate_mean", "Migration")])
    add_fig(doc, FIG / "fig_qtbias_fixed_checkpoint_robustness.png", "Fig. 7. Fixed-checkpoint robustness.")

    doc.add_heading("6.7 KAIST N=40 高负载压力测试", 2)
    para("N=40 是额外 high-load stress test。PriceEdge-QT-Bias 在高负载下仍避免 PPO-DSR-A/PPO-MV-A 接近 10s CVaR 饱和的情况，但应解释为相对抗压能力，而非持续过载下的绝对服务保证。")
    para("表 5. KAIST N=40 pressure test。", WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, n40, [("method", "方法"), ("distribution", "分布"), ("runs", "runs"), ("average_latency_mean", "Avg latency"), ("cvar99_mean", "CVaR99"), ("queue_oscillation_mean", "Queue osc."), ("portfolio_migration_rate_mean", "Migration")])
    add_fig(doc, FIG / "fig_qtbias_n40_stress.png", "Fig. 8. KAIST N=40 stress test.")

    doc.add_heading("7 讨论", 1)
    para("当前证据支持三点：共享滞后反馈下的 portfolio 卸载是多用户耦合问题；QT price-to-logit bias 与正确 price-server 对齐有助于降低 CVaR99；edge-wise actor 比 flat actor 更适合绑定服务器特征、价格和动作 logit。同时，stale margin 与 tail state Z 在当前结果中不是核心增益来源。")
    para("未来需要补充 random time-varying delay、per-server information age、missed broadcast、asynchronous user updates、capacity shock 和 hidden background load，以检验逐边结构在更复杂信息不一致条件下的价值。")

    doc.add_heading("8 结论", 1)
    para("本文提出 PriceEdge-QT-Bias MAPPO，用逐边队列—时延价格和 edge-wise shared actor 处理共享滞后反馈下的多用户协作式 portfolio 卸载。实验显示该方法在统一 adapted-baseline 接口下显著降低平均时延和 CVaR99，并在固定 checkpoint 跨环境测试与 N=40 压力测试中保持相对优势。")

    doc.add_heading("参考文献", 1)
    refs = markdown.split("## 参考文献", 1)[1].strip().splitlines()
    for ref in refs:
        if ref.strip():
            para(ref.replace("  ", " "))
    doc.save(DOCX_OUT)


def main() -> None:
    create_architecture_figure()
    md = build_markdown()
    MD_OUT.write_text(md, encoding="utf-8")
    build_docx(md)
    print(MD_OUT)
    print(DOCX_OUT)
    print(ARCH_FIG)


if __name__ == "__main__":
    main()
