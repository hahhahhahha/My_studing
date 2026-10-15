"""Build an editable, readable QLA paper DOCX from audited text and existing evidence.

The body uses the code-aligned v7 system/method text, while the experiment
section is copied from the evidence-rich v6 draft.  Values known to depend on
the pending Full-noZ/shared-schedule/JSEQ repair remain explicitly provisional.
"""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MAIN = ROOT / "paper_draft_qla_cn_v7_main_editable.md"
APPENDIX = ROOT / "paper_draft_qla_cn_v7_appendix_editable.md"
EVIDENCE = ROOT / "paper_draft_priceedge_qtbias_cn_v6_literature_audited.md"
REFERENCE = ROOT / "paper_draft_facr_cn_v9_evidence_checked.docx"
OUTPUT = ROOT / "paper_draft_qla_cn_v7_readable_provisional.docx"


def section(text: str, start: str, end: str | None = None) -> str:
    left = text.index(start)
    if end is None:
        return text[left:]
    right = text.index(end, left)
    return text[left:right]


def build_markdown() -> str:
    main = MAIN.read_text(encoding="utf-8")
    evidence = EVIDENCE.read_text(encoding="utf-8")
    appendix = APPENDIX.read_text(encoding="utf-8")

    prefix = main[: main.index("## 6 实验结果")]
    experiments = section(evidence, "## 6 实验结果", "## 7 讨论与局限")
    experiments = experiments.replace(
        "```text\nD_stale(M) = (M_{δ=8} - M_{δ=0}) / M_{δ=0}.\n```",
        "$$\nD_{\\mathrm{stale}}(M)=\\frac{M_{\\delta=8}-M_{\\delta=0}}{M_{\\delta=0}}.\n$$",
    )
    experiments = experiments.replace("`Q_j^{\\max}`", "$Q_j^{\\max}$")
    experiments = experiments.replace("`P(T≥10s)`", "$P(T\\ge 10\\,\\mathrm{s})$")
    suffix = main[main.index("## 7 讨论") :]

    # The earlier scenario placeholder is removed from this readable draft.
    # A dedicated publication illustration will be inserted after it is drawn.
    prefix = re.sub(
        r"> \*\*Fig\. 1 占位：.*?\*\*\s*",
        "",
        prefix,
        flags=re.DOTALL,
    )
    architecture = (
        "![Fig. 1. QLA-MAPPO 的代码对齐架构：共享滞后摘要经逐边 QL 信号进入参数共享的 "
        "edge-wise Gaussian Actor，并通过联合队列演化和 CTDE PPO 训练。]"
        "(X:/Yi-one/workspace/MPT-CVaR-MAPPO/results_priceedge_qtbias_paper_ready/figures/"
        "fig_priceedge_qtbias_architecture.png)"
    )
    prefix = re.sub(
        r"> \*\*Fig\. 2 占位：.*?\*\*\s*",
        architecture + "\n\n",
        prefix,
        flags=re.DOTALL,
    )

    replacements = {
        "<<DATASET/SEED/AGE PROTOCOL>>": "KAIST 的 5 个训练种子 × 5 个 nominal feedback delay 兼容性协议（暂定）",
        "<<AVG>>": "0.896",
        "<<CVAR99>>": "3.241",
        "<<AVG_EFFECT>>": "29.7",
        "<<CVAR_EFFECT>>": "52.7",
        "<<CI_AVG>>": "[−0.448, −0.300] s（相对 PPO-MV-A）",
        "<<CI_CVAR>>": "[−3.730, −3.453] s（相对 QECO-A）",
        "<<QUEUE-BOUNDARY-PROTOCOL>>": "[队列守恒修复后冻结；当前兼容性结果使用有限队列裁剪实现]",
        "<<AGE-INPUT-PROTOCOL>>": "[source-age 输入协议修复后冻结；当前兼容性结果使用 nominal delay proxy]",
    }
    body = prefix + experiments + suffix
    for old, new in replacements.items():
        body = body.replace(old, new)
    body = re.sub(r"<<[^>]*>>", "[待协议修复后更新]", body)

    notice = (
        "> **可读暂定稿。** 正文已回填现有 v6 图表和数据。凡依赖 Full-noZ、共享评估序列、"
        "JSEQ-A 修复或队列守恒协议的数值均属于兼容性证据，最终投稿前必须由修复后的正式结果替换。\n\n"
    )
    first_break = body.find("\n") + 1
    body = body[:first_break] + "\n" + notice + body[first_break:]

    appendix = re.sub(r"<<[^>]*>>", "[待协议修复后更新]", appendix)
    appendix = re.sub(r"^# .+?\n", "", appendix, count=1)
    return body.rstrip() + "\n\n\\newpage\n\n# 补充材料\n\n" + appendix.strip() + "\n"


def polish_docx(path: Path) -> None:
    doc = Document(path)
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(0.72)
        sec.bottom_margin = Inches(0.72)
        sec.left_margin = Inches(0.72)
        sec.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    for paragraph in doc.paragraphs:
        xml = paragraph._p.xml
        text = paragraph.text.strip()
        if "<w:drawing" in xml or "<m:oMath" in xml or "<m:oMathPara" in xml:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text.startswith(("Fig.", "图 ", "表 ", "Table ")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(9)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)

    doc.save(path)


def main() -> None:
    merged = build_markdown()
    # Pandoc's texmath reader rejects the legacy TeX ``\rm`` switch.  Convert
    # it mechanically to ``\mathrm{...}`` so equations become editable OMML
    # objects instead of literal TeX text in Word.
    merged = re.sub(r"\\rm\s+([A-Za-z]+)", r"\\mathrm{\1}", merged)
    with tempfile.TemporaryDirectory(prefix="qla_docx_") as tmp:
        source = Path(tmp) / "paper.md"
        source.write_text(merged, encoding="utf-8")
        cmd = [
            "pandoc",
            str(source),
            "--from=markdown+tex_math_dollars+raw_tex",
            "--to=docx",
            "--standalone",
            "--toc",
            "--toc-depth=3",
            f"--resource-path={ROOT}",
            f"--reference-doc={REFERENCE}",
            "--metadata=lang:zh-CN",
            "--output",
            str(OUTPUT),
        ]
        subprocess.run(cmd, cwd=ROOT, check=True)
    polish_docx(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
