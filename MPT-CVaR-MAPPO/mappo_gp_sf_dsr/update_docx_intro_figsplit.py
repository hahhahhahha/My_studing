"""Insert split introduction figures into paper_draft_facr_cn_v7.docx.

The source DOCX is preserved. Existing "Fig. n" references are shifted by +2
before inserting the new Fig. 1 and Fig. 2.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "paper_draft_facr_cn_v7.docx"
OUT = ROOT / "paper_draft_facr_cn_v7_fig1split.docx"
FIG1 = ROOT / "figures" / "fig_intro_stale_feedback_mechanism.png"
FIG2 = ROOT / "figures" / "fig_intro_scoresoftmax_delay_evidence.png"


FIG_RE = re.compile(r"\bFig\.\s*(\d+)([a-z]?)")


def shift_fig_refs(text: str, offset: int = 2) -> str:
    def repl(match: re.Match[str]) -> str:
        return f"Fig. {int(match.group(1)) + offset}{match.group(2)}"

    return FIG_RE.sub(repl, text)


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


def set_run_font(run, size_pt: float = 10.0, name: str = "Times New Roman") -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "SimSun")


def add_centered_image_after(anchor: Paragraph, image_path: Path, width_in: float = 6.4) -> Paragraph:
    p = insert_paragraph_after(anchor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))
    return p


def add_caption_after(anchor: Paragraph, caption: str) -> Paragraph:
    p = insert_paragraph_after(anchor, style="Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    set_run_font(run, size_pt=9.5)
    return p


def add_body_after(anchor: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(anchor, style="Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size_pt=10.5)
    return p


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    for fig in (FIG1, FIG2):
        if not fig.exists():
            raise FileNotFoundError(fig)

    doc = Document(str(SRC))

    # Shift all pre-existing figure references/captions by +2.
    for para in doc.paragraphs:
        for run in para.runs:
            if "Fig." in run.text:
                run.text = shift_fig_refs(run.text, 2)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if "Fig." in run.text:
                            run.text = shift_fig_refs(run.text, 2)

    # Insert after the paragraph ending the problem formulation in the introduction,
    # just before "为此，本文提出 FACR-MAPPO".
    anchor = None
    for idx, para in enumerate(doc.paragraphs):
        if para.text.startswith("为此，本文提出 FACR-MAPPO"):
            anchor = doc.paragraphs[idx - 1]
            break
    if anchor is None:
        raise RuntimeError("Could not locate the introduction insertion point.")

    p = add_body_after(
        anchor,
        "如 Fig. 1 所示，本文关注的场景机制是：用户只能读取滞后的边缘状态摘要，"
        "当多个用户同时追逐同一份历史低负载指示时，局部误判会被放大为队列尖峰和尾部时延。"
        "Fig. 2 进一步在 KAIST、N=30、score-softmax 压力策略下给出实证证据，"
        "说明 feedback age 增大时 CVaR99 和队列振荡均显著恶化。",
    )
    p = add_centered_image_after(p, FIG1)
    p = add_caption_after(
        p,
        "Fig. 1. 滞后边缘状态反馈诱发系统性拥挤的机制示意。"
        "陈旧摘要可能保留一定统计相关性，但当多用户同步响应同一份历史低负载指示时，"
        "相关动作会放大瞬时队列尖峰和 tail-latency events。",
    )
    p = add_centered_image_after(p, FIG2)
    p = add_caption_after(
        p,
        "Fig. 2. KAIST、N=30、score-softmax 压力策略下的 feedback-age 敏感性。"
        "随着 δ 从 0 增至 8，CVaR99 由 5.81 s 上升至 10.04 s，"
        "queue oscillation 由 0.84 上升至 3.77，表明滞后反馈会显著放大尾部时延和队列振荡。",
    )

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
